"""Crea el ensayo Word conservando el paquete de la plantilla de IPC2.

Requiere Python, lxml, python-docx y Pillow. Exportar el PDF con
Exportar_documentacion.ps1 requiere Microsoft Word en Windows.
Estas herramientas no son necesarias para ejecutar el simulador SAT/DTE.
"""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from lxml import etree as ET

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "plantillas" / "Ensayo_IPC2.docx"
OUTPUT = ROOT / "Documentacion.docx"
TEMPLATE_SHA = "79217801f046dd99f238fd33cfbf1501bab4f519a9c369dc389afa72b63204eb"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def tag(name):
    prefix, local = name.split(":")
    return "{" + NS[prefix] + "}" + local


def element(name, **attributes):
    node = ET.Element(tag(name))
    for key, value in attributes.items():
        node.set(tag(key) if ":" in key else key, str(value))
    return node


assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == TEMPLATE_SHA
content = json.loads((ROOT / "ensayo_contenido.json").read_text(encoding="utf-8"))
assert 150 <= len(content["summary"].split()) <= 175
assert len(content["introduction"].split()) <= 150
assert len(content["keywords"]) <= 5 and len(content["references"]) <= 5
with ZipFile(TEMPLATE) as source:
    parts = {item.filename: source.read(item.filename) for item in source.infolist()}
document = ET.fromstring(parts["word/document.xml"])
body = document.find("w:body", NS)
source_paragraphs = body.findall("w:p", NS)
original_end = deepcopy(body.find("w:sectPr", NS))
styles = {role: source_paragraphs[index] for role, index in {
    "heading": 46, "summary_heading": 3, "keywords_heading": 14,
    "english_keywords_heading": 38, "body": 58, "summary": 5, "english": 26,
    "english_heading": 24, "subheading": 62, "caption": 69,
    "source": 70, "reference": 101,
}.items()}


def paragraph(text, role="body", *, bold=False, italic=None):
    pattern = styles[role]
    node = element("w:p")
    properties = pattern.find("w:pPr", NS)
    properties = deepcopy(properties) if properties is not None else element("w:pPr")
    for child in list(properties):
        if child.tag in {tag("w:sectPr"), tag("w:rPr")}:
            properties.remove(child)
    node.append(properties)
    if role.endswith("heading") or bold:
        properties.append(element("w:keepNext"))
    if role == "reference":
        old = properties.find("w:ind", NS)
        if old is not None:
            properties.remove(old)
        properties.append(element("w:ind", **{"w:left": "720", "w:hanging": "720"}))
        alignment = properties.find("w:jc", NS)
        if alignment is not None:
            properties.remove(alignment)
        properties.append(element("w:jc", **{"w:val": "left"}))
    run = element("w:r")
    original = pattern.find("w:r/w:rPr", NS)
    run_properties = deepcopy(original) if original is not None else element("w:rPr")
    if bold and run_properties.find("w:b", NS) is None:
        run_properties.append(element("w:b"))
    if italic is not None:
        old = run_properties.find("w:i", NS)
        if old is not None:
            run_properties.remove(old)
        run_properties.append(element("w:i", **{"w:val": "1" if italic else "0"}))
    language = run_properties.find("w:lang", NS)
    if language is None:
        language = element("w:lang")
        run_properties.append(language)
    language.set(tag("w:val"), "en-US" if role.startswith("english") else "es-GT")
    run.append(run_properties)
    text_node = element("w:t")
    text_node.text = text
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(text_node)
    node.append(run)
    return node


def break_paragraph(kind):
    node = element("w:p")
    properties = element("w:pPr")
    properties.append(element("w:spacing", **{"w:before": "0", "w:after": "0", "w:line": "20", "w:lineRule": "exact"}))
    node.append(properties)
    run = element("w:r")
    run.append(element("w:br", **{"w:type": kind}))
    node.append(run)
    return node


# Se conservan los elementos originales del bloque de título y su sección.
prefix = [deepcopy(child) for child in list(body)[:4]]
slots = [p for p in prefix[2].findall(".//w:p", NS) if "".join(p.xpath(".//w:t/text()", namespaces=NS)).strip()]
for slot, replacement in zip(slots, [content["title"], content["author"]]):
    texts = slot.findall(".//w:t", NS)
    texts[0].text = replacement
    for text in texts[1:]:
        text.text = ""
for child in list(body):
    body.remove(child)
for child in prefix:
    body.append(child)
body.append(paragraph("Resumen", "summary_heading"))
body.append(paragraph(content["summary"], "summary"))
body.append(paragraph("Palabras clave", "keywords_heading"))
body.append(paragraph("; ".join(content["keywords"])))
body.append(break_paragraph("column"))
body.append(paragraph("Abstract", "english_heading"))
body.append(paragraph(content["abstract"], "english", italic=True))
body.append(paragraph("Keywords", "english_keywords_heading"))
body.append(paragraph("; ".join(content["english_keywords"]), "english", italic=True))
body.append(break_paragraph("page"))
body.append(paragraph("Introducción", "heading"))
body.append(paragraph(content["introduction"]))
body.append(paragraph("Desarrollo del tema", "heading"))
for section in content["sections"]:
    body.append(paragraph(section["title"], "subheading", bold=True))
    for text in section["paragraphs"]:
        body.append(paragraph(text))
body.append(paragraph("Conclusiones", "heading"))
for text in content["conclusions"]:
    body.append(paragraph(text))
body.append(paragraph("Referencias bibliográficas", "heading"))
for reference in content["references"]:
    before = f"{reference['author']} ({reference['year']}). "
    after = f". {reference['publisher']}." + (" " + reference["url"] if reference["url"] else "")
    node = paragraph(before, "reference")
    pattern_run = node.find("w:r", NS)
    for text, is_italic in [(reference["title"], True), (after, False)]:
        run = deepcopy(pattern_run)
        run.find("w:t", NS).text = text
        properties = run.find("w:rPr", NS)
        old = properties.find("w:i", NS)
        if old is not None:
            properties.remove(old)
        properties.append(element("w:i", **{"w:val": "1" if is_italic else "0"}))
        node.append(run)
    body.append(node)

# Una nueva sección permite anexos de una columna sin cambiar el ensayo.
essay_section = deepcopy(original_end)
section_break = element("w:p")
properties = element("w:pPr")
properties.append(essay_section)
section_break.append(properties)
body.append(section_break)
appendix_section = deepcopy(original_end)
section_type = appendix_section.find("w:type", NS)
if section_type is None:
    section_type = element("w:type")
    appendix_section.insert(0, section_type)
section_type.set(tag("w:val"), "nextPage")
columns = appendix_section.find("w:cols", NS)
for child in list(columns):
    columns.remove(child)
columns.attrib.clear()
columns.set(tag("w:num"), "1")
columns.set(tag("w:space"), "720")

relationships = ET.fromstring(parts["word/_rels/document.xml.rels"])
relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
for index, appendix in enumerate(content["appendices"], 1):
    if index > 1:
        body.append(break_paragraph("page"))
    body.append(paragraph(appendix["title"], "heading"))
    image_path = ROOT / "imagenes" / appendix["image"]
    # python-docx construye únicamente el fragmento de imagen; el documento
    # final sigue siendo el paquete original, con sus partes preservadas.
    image_document = Document()
    picture_paragraph = image_document.add_paragraph()
    shape = picture_paragraph.add_run().add_picture(str(image_path), width=Inches(5.9 if index < 4 else 6.7))
    shape._inline.docPr.set("id", str(3000 + index))
    shape._inline.docPr.set("descr", appendix["caption"])
    picture_node = deepcopy(picture_paragraph._p)
    picture_properties = element("w:pPr")
    picture_properties.append(element("w:jc", **{"w:val": "center"}))
    picture_properties.append(element("w:spacing", **{"w:before": "0", "w:after": "60"}))
    picture_node.insert(0, picture_properties)
    relation_id = f"rIdEnsayoFigura{index}"
    for blip in picture_node.xpath('.//*[local-name()="blip"]'):
        blip.set(tag("r:embed"), relation_id)
    media_name = f"ensayo_figura_{index}.png"
    parts["word/media/" + media_name] = image_path.read_bytes()
    relationship = ET.SubElement(relationships, "{" + relationship_ns + "}Relationship")
    relationship.set("Id", relation_id)
    relationship.set("Type", NS["r"] + "/image")
    relationship.set("Target", "media/" + media_name)
    body.append(picture_node)
    body.append(paragraph(appendix["caption"], "caption", italic=False))
    body.append(paragraph("Fuente: elaboración propia.", "source"))
body.append(appendix_section)

content_types = ET.fromstring(parts["[Content_Types].xml"])
if not any(n.get("Extension") == "png" for n in content_types):
    item = ET.SubElement(content_types, "{" + content_types.nsmap[None] + "}Default")
    item.set("Extension", "png")
    item.set("ContentType", "image/png")
for name, node in [("word/document.xml", document), ("word/_rels/document.xml.rels", relationships), ("[Content_Types].xml", content_types)]:
    parts[name] = ET.tostring(node, xml_declaration=True, encoding="UTF-8", standalone=True)
core = ET.fromstring(parts["docProps/core.xml"])
for key, value in {
    "{http://purl.org/dc/elements/1.1/}title": content["title"],
    "{http://purl.org/dc/elements/1.1/}creator": "William René Toledo Corado",
    "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastModifiedBy": "William René Toledo Corado",
}.items():
    node = core.find(key)
    if node is None:
        node = ET.SubElement(core, key)
    node.text = value
modified = core.find("{http://purl.org/dc/terms/}modified")
if modified is not None:
    modified.text = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
parts["docProps/core.xml"] = ET.tostring(core, xml_declaration=True, encoding="UTF-8", standalone=True)
with ZipFile(OUTPUT, "w", compression=ZIP_DEFLATED) as target:
    for name, data in parts.items():
        target.writestr(name, data)
editable = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "docProps/core.xml", "docProps/app.xml"}
with ZipFile(TEMPLATE) as reference, ZipFile(OUTPUT) as final:
    for name in reference.namelist():
        if name not in editable:
            assert reference.read(name) == final.read(name), "Parte preservada modificada: " + name
print(OUTPUT)
print("Resumen:", len(content["summary"].split()), "palabras; introducción:", len(content["introduction"].split()), "palabras.")
